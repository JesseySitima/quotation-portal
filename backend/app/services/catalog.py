from io import BytesIO

from fastapi import HTTPException, UploadFile
from openpyxl import load_workbook

from app.db.client import supabase


def upload_catalog_service(file: UploadFile):
    if not file.filename:
        raise HTTPException(
            status_code=400,
            detail="No file selected",
        )

    if not file.filename.lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(
            status_code=400,
            detail="Please upload an Excel file (.xlsx or .xlsm)",
        )

    try:
        file_bytes = file.file.read()

        workbook = load_workbook(
            filename=BytesIO(file_bytes),
            read_only=True,
            data_only=True,
        )

        worksheet = workbook.active

        rows = list(
            worksheet.iter_rows(
                values_only=True,
            )
        )

        workbook.close()

    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail=f"Could not read Excel file: {str(exc)}",
        )

    if not rows:
        raise HTTPException(
            status_code=400,
            detail="The Excel file is empty",
        )

    headers = [
        str(value).strip() if value is not None else ""
        for value in rows[0]
    ]

    required_headers = [
        "No.",
        "Description",
        "Storage Place",
        "Base Unit of Measure",
    ]

    if headers[:4] != required_headers:
        raise HTTPException(
            status_code=400,
            detail=(
                "Invalid catalog format. "
                "Expected columns: "
                "No., Description, Storage Place, "
                "Base Unit of Measure"
            ),
        )

    # Load active categories once.
    category_response = (
        supabase
        .table("categories")
        .select("id, name, slug")
        .eq("is_active", True)
        .execute()
    )

    categories = category_response.data or []

    category_map = {
        category["name"].strip().upper(): category["id"]
        for category in categories
    }

    # Load existing products once.
    product_response = (
        supabase
        .table("products")
        .select("id, sku, name, unit, category_id")
        .execute()
    )

    existing_products = product_response.data or []

    product_map = {
        str(product["sku"]).strip(): product
        for product in existing_products
    }

    errors = []
    products_to_upsert = []

    rows_processed = 0
    new_products = 0
    existing_products_count = 0

    for row_number, row in enumerate(rows[1:], start=2):
        if not row:
            continue

        sku = row[0]
        name = row[1]
        storage_place = row[2]
        unit = row[3]

        # Ignore completely empty rows.
        if all(
            value is None or str(value).strip() == ""
            for value in row[:4]
        ):
            continue

        rows_processed += 1

        if sku is None or str(sku).strip() == "":
            errors.append(
                f"Row {row_number}: No. is required"
            )
            continue

        if name is None or str(name).strip() == "":
            errors.append(
                f"Row {row_number}: Description is required"
            )
            continue

        if storage_place is None or str(storage_place).strip() == "":
            errors.append(
                f"Row {row_number}: Storage Place is required"
            )
            continue

        if unit is None or str(unit).strip() == "":
            errors.append(
                f"Row {row_number}: Base Unit of Measure is required"
            )
            continue

        sku = str(sku).strip()
        name = str(name).strip()
        storage_place = str(storage_place).strip()
        unit = str(unit).strip()

        category_id = category_map.get(
            storage_place.upper()
        )

        if not category_id:
            errors.append(
                f"Row {row_number}: Unknown Storage Place "
                f"'{storage_place}'"
            )
            continue

        existing_product = product_map.get(sku)

        if existing_product:
            existing_products_count += 1
        else:
            new_products += 1

        products_to_upsert.append(
            {
                "sku": sku,
                "name": name,
                "category_id": category_id,
                "unit": unit,
                "description": None,
                "is_available": True,
            }
        )

    if errors:
        raise HTTPException(
            status_code=400,
            detail={
                "message": "Catalog validation failed",
                "rows_processed": rows_processed,
                "errors": errors,
            },
        )

    if not products_to_upsert:
        raise HTTPException(
            status_code=400,
            detail="No valid catalog products found",
        )

    try:
        response = (
            supabase
            .table("products")
            .upsert(
                products_to_upsert,
                on_conflict="sku",
            )
            .execute()
        )

    except Exception as exc:
        raise HTTPException(
            status_code=500,
            detail=f"Failed to update product catalog: {str(exc)}",
        )

    return {
        "message": "Catalog uploaded successfully",
        "filename": file.filename,
        "rows_processed": rows_processed,
        "new_products": new_products,
        "existing_products": existing_products_count,
        "products_updated": len(response.data or []),
    }